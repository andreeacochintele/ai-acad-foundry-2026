"""Extract plain text from a document a user attaches to one conversation.

Deliberately narrow: four formats (txt, md, pdf, docx), each read with a
small, well-known library — no OCR, no scanned-PDF fallback, no layout
preservation. That's the honest version of "read my document," matching
services/web.py's own stance on scraping: report what this could not do
rather than pretend it handled everything.

This never touches Qdrant or the persistent knowledge base — the extracted
text is only ever included in the prompt for the question(s) it was
attached to (see AskRequest.attached_document), never ingested.
"""
from __future__ import annotations

import io
import re
from dataclasses import dataclass, field

SUPPORTED_EXTENSIONS = {"txt", "md", "pdf", "docx"}


class UnsupportedDocumentType(Exception):
    """The file extension isn't one of SUPPORTED_EXTENSIONS."""


@dataclass
class DocumentExtractResult:
    filename: str
    file_type: str
    text: str
    chars: int
    approx_tokens: int
    warnings: list[str] = field(default_factory=list)


def _extract_txt(content: bytes) -> str:
    return content.decode("utf-8", errors="replace")


def _extract_pdf(content: bytes) -> tuple[str, list[str]]:
    from pypdf import PdfReader

    warnings: list[str] = []
    reader = PdfReader(io.BytesIO(content))
    if reader.is_encrypted:
        try:
            reader.decrypt("")   # only ever tries an empty password, never guesses one
        except Exception:
            return "", ["This PDF is password-protected — cannot extract text from it."]
    pages = []
    for i, page in enumerate(reader.pages):
        try:
            pages.append(page.extract_text() or "")
        except Exception as e:
            warnings.append(f"Page {i + 1} failed to extract: {e}")
    text = "\n\n".join(p for p in pages if p.strip())
    if not text.strip():
        warnings.append(
            "No extractable text found — this is likely a scanned/image-only PDF. "
            "A plain text extractor cannot read that; it would need OCR."
        )
    return text, warnings


def _extract_docx(content: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(content))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n\n".join(parts)


def extract_text(filename: str, content: bytes, max_chars: int = 20000) -> DocumentExtractResult:
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext not in SUPPORTED_EXTENSIONS:
        raise UnsupportedDocumentType(
            f"'.{ext or '(none)'}' is not supported — attach one of: "
            f"{', '.join(sorted(SUPPORTED_EXTENSIONS))}."
        )

    warnings: list[str] = []
    if ext in ("txt", "md"):
        text = _extract_txt(content)
    elif ext == "pdf":
        text, pdf_warnings = _extract_pdf(content)
        warnings.extend(pdf_warnings)
    else:  # docx
        text = _extract_docx(content)

    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if len(text) > max_chars:
        text = text[:max_chars]
        warnings.append(f"Truncated to {max_chars} characters.")

    return DocumentExtractResult(
        filename=filename,
        file_type=ext,
        text=text,
        chars=len(text),
        approx_tokens=max(1, round(len(text) / 4)),
        warnings=warnings,
    )
